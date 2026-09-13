from __future__ import annotations

import re
from docx import Document
from core.schema import Investigation
from reports.common import NOTICE, entity_rows, tool_rows, text_sections


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '\ufffd', str(value))


def generate_docx(inv: Investigation, output_path: str) -> str:
    doc = Document()
    doc.add_heading('OSINT-X AI Investigation Report', 0)
    for text in (f'ID: {inv.id}', f'Target: {inv.target_type} = {inv.target_value}', f'Started: {inv.started_at}', f'Completed: {inv.completed_at}', NOTICE):
        doc.add_paragraph(clean(text))
    for title, text in text_sections(inv):
        doc.add_heading(title, 1)
        doc.add_paragraph(clean(text))
    for title, entities in inv.entity_groups():
        doc.add_heading(title, 1)
        if not entities:
            doc.add_paragraph('None.')
        for entity in entities:
            doc.add_heading(clean(f'{entity.type.value}: {entity.value}'), 2)
            for text in (f'Status: {entity.status.value}; source: {entity.source}', f'Evidence: {entity.evidence}', f'Confidence basis: {entity.confidence_basis}', f'Metadata: {entity.metadata}'):
                doc.add_paragraph(clean(text))
    doc.add_heading('Source results / errors / unavailable sources', 1)
    for row in tool_rows(inv):
        doc.add_paragraph(clean(' | '.join(row)))
    doc.add_heading('Search suggestions — UNVERIFIED', 1)
    for s in inv.suggestions:
        doc.add_paragraph(clean(s['label'] + ': ' + s['url']))
    for title, texts in [('Heuristic exposure assessment', [str(inv.risk)]), ('Timeline', [e.timestamp + ': ' + e.description for e in inv.timeline]), ('AI-selected follow-up actions', inv.ai_recommendations), ('Warnings / limitations', inv.warnings)]:
        doc.add_heading(title, 1)
        for text in texts:
            doc.add_paragraph(clean(text))
    doc.save(output_path)
    return output_path
